import docx
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

doc = docx.Document()

# Add a title
title = doc.add_heading('4.3 Architectural Design', level=1)
title.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT

doc.add_paragraph('This section presents the architectural design for the VALORA Property Health Intelligence Platform, detailing the software, hardware, and network components, along with the core class diagram.')

# 4.3.1 Software Architectural Design
doc.add_heading('4.3.1 Software Architectural Design', level=2)
p = doc.add_paragraph('VALORA employs a modern, layered Client-Server architecture based on the Model-View-Controller (MVC) pattern and RESTful principles, optimized for scalability and modularity. The primary components include:')
ul1 = doc.add_paragraph(style='List Bullet')
ul1.add_run('Frontend (Presentation Layer): ').bold = True
ul1.add_run('A Single Page Application (SPA) built with React.js. It interacts with the backend via RESTful APIs to provide dynamic user interfaces for Residents, Investors, HOA Admins, and Service Providers.')

ul2 = doc.add_paragraph(style='List Bullet')
ul2.add_run('Backend (Application Layer): ').bold = True
ul2.add_run('A Python-based FastAPI application that handles business logic, API requests, and authentication. It manages the core complaint lifecycle and work order state machine.')

ul3 = doc.add_paragraph(style='List Bullet')
ul3.add_run('Property Health Engine: ').bold = True
ul3.add_run('A dedicated AI and analytics module within the backend that calculates the Living Score and Property Health Score using operational community data (complaints, work orders, ratings) and integrates this with valuation models.')

ul4 = doc.add_paragraph(style='List Bullet')
ul4.add_run('Data Layer: ').bold = True
ul4.add_run('A PostgreSQL relational database managed via SQLAlchemy ORM and Alembic for schema migrations. It persistently stores user profiles, property records, complaints, maintenance work orders, and AI analytics.')

# 4.3.2 Hardware Architectural Design
doc.add_heading('4.3.2 Hardware Architectural Design', level=2)
doc.add_paragraph('The system is deployed using a cloud-native infrastructure leveraging Docker containerization to ensure high availability and ease of deployment.')
hl1 = doc.add_paragraph(style='List Bullet')
hl1.add_run('Client Devices: ').bold = True
hl1.add_run('Users access VALORA via web browsers on standard desktops, laptops, tablets, or mobile devices. No specialized hardware is required on the client side.')
hl2 = doc.add_paragraph(style='List Bullet')
hl2.add_run('Web/Application Servers: ').bold = True
hl2.add_run('Dockerized containers hosting the FastAPI backend and React frontend, deployed on scalable cloud compute instances (e.g., AWS EC2, AWS ECS, or Azure Virtual Machines) capable of auto-scaling based on traffic.')
hl3 = doc.add_paragraph(style='List Bullet')
hl3.add_run('Database Server: ').bold = True
hl3.add_run('A managed relational database instance (e.g., AWS RDS for PostgreSQL) configured with automated backups, Multi-AZ high availability, and failover capabilities.')
hl4 = doc.add_paragraph(style='List Bullet')
hl4.add_run('Storage: ').bold = True
hl4.add_run('Cloud object storage (e.g., AWS S3) for storing user-uploaded media attachments related to complaints (e.g., photos of maintenance issues) and static assets.')

# 4.3.3 Network Architectural Design
doc.add_heading('4.3.3 Network Architectural Design', level=2)
doc.add_paragraph('The network architecture is designed to be secure, resilient, and performant, utilizing standard internet protocols.')
nl1 = doc.add_paragraph(style='List Bullet')
nl1.add_run('Internet & Security: ').bold = True
nl1.add_run('Clients connect over the public internet using HTTPS (TLS 1.2/1.3) to ensure data encryption in transit. JWT tokens are used for stateless API authentication.')
nl2 = doc.add_paragraph(style='List Bullet')
nl2.add_run('Load Balancer / API Gateway: ').bold = True
nl2.add_run('A cloud load balancer sits at the edge of the network, distributing incoming HTTP/HTTPS traffic evenly across the frontend and backend server instances, and providing SSL termination.')
nl3 = doc.add_paragraph(style='List Bullet')
nl3.add_run('Virtual Private Cloud (VPC): ').bold = True
nl3.add_run('The platform resides within a private VPC. The frontend and load balancer are exposed via the public subnet. The backend APIs and the database are strictly isolated in private subnets, accessible only through the load balancer and backend application servers respectively.')
nl4 = doc.add_paragraph(style='List Bullet')
nl4.add_run('API Communication: ').bold = True
nl4.add_run('The React frontend communicates with the FastAPI backend exclusively via secure RESTful JSON APIs.')

# 4.3.4 Class Diagram
doc.add_heading('4.3.4 Class Diagram', level=2)
doc.add_paragraph('The class diagram models the core domain entities of the VALORA system and their structural relationships.')

doc.add_heading('Core Classes & Attributes:', level=3)
cl1 = doc.add_paragraph(style='List Bullet')
cl1.add_run('User: ').bold = True
cl1.add_run('user_id, name, email, role_id, password_hash, created_at')
cl2 = doc.add_paragraph(style='List Bullet')
cl2.add_run('Role: ').bold = True
cl2.add_run('role_id, role_name (e.g., Resident, Investor, HOA Admin, Provider)')
cl3 = doc.add_paragraph(style='List Bullet')
cl3.add_run('Property: ').bold = True
cl3.add_run('property_id, address, unit_number, complex_name, size, bedrooms, bathrooms')
cl4 = doc.add_paragraph(style='List Bullet')
cl4.add_run('Complaint: ').bold = True
cl4.add_run('complaint_id, resident_id, property_id, category, status, description, timestamp')
cl5 = doc.add_paragraph(style='List Bullet')
cl5.add_run('MaintenanceWorkOrder: ').bold = True
cl5.add_run('order_id, complaint_id, provider_id, status, cost, resident_rating, resolution_time')
cl6 = doc.add_paragraph(style='List Bullet')
cl6.add_run('ServiceProvider: ').bold = True
cl6.add_run('provider_id, company_name, specialization, average_rating, is_verified')
cl7 = doc.add_paragraph(style='List Bullet')
cl7.add_run('AIAnalytic: ').bold = True
cl7.add_run('analytic_id, property_id, living_score, predicted_value, calculation_date')

doc.add_heading('Relationships:', level=3)
r1 = doc.add_paragraph(style='List Bullet')
r1.add_run('User to Role (Many-to-One): ').bold = True
r1.add_run('Many users can have one specific role.')
r2 = doc.add_paragraph(style='List Bullet')
r2.add_run('User to Property (One-to-Many): ').bold = True
r2.add_run('A user (Resident/Investor) can be associated with one or more properties.')
r3 = doc.add_paragraph(style='List Bullet')
r3.add_run('User to Complaint (One-to-Many): ').bold = True
r3.add_run('A Resident logs one or more complaints.')
r4 = doc.add_paragraph(style='List Bullet')
r4.add_run('Property to Complaint (One-to-Many): ').bold = True
r4.add_run('A property can have multiple complaints associated with it.')
r5 = doc.add_paragraph(style='List Bullet')
r5.add_run('Complaint to MaintenanceWorkOrder (One-to-One): ').bold = True
r5.add_run('Each valid complaint generates a distinct maintenance work order.')
r6 = doc.add_paragraph(style='List Bullet')
r6.add_run('ServiceProvider to MaintenanceWorkOrder (One-to-Many): ').bold = True
r6.add_run('A verified service provider can handle multiple work orders.')
r7 = doc.add_paragraph(style='List Bullet')
r7.add_run('Property to AIAnalytic (One-to-Many): ').bold = True
r7.add_run('A property has multiple analytics records generated over time to track its Living Score and value.')

doc.save(r'C:\Users\garah\PycharmProjects\Group8Project\docs\Valora_Architectural_Design.docx')
print("File created successfully.")
